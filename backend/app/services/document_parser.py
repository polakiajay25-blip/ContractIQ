
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def extract_text(file_path: str) -> dict:
    """
    Extract text from a PDF or DOCX file.

    PDF results include page numbers.
    DOCX results include paragraph numbers.
    """
    path = Path(file_path)

    if not path.exists() or not path.is_file():
        raise FileNotFoundError("The uploaded document was not found.")

    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Only PDF and DOCX files are supported.")

    sections = []

    if extension == ".pdf":
        with fitz.open(path) as pdf:
            for page_number, page in enumerate(pdf, start=1):
                text = page.get_text("text").strip()

                if text:
                    sections.append({
                        "location": f"Page {page_number}",
                        "page": page_number,
                        "text": text,
                    })

    elif extension == ".docx":
        document = Document(path)

        for paragraph_number, paragraph in enumerate(
            document.paragraphs, start=1
        ):
            text = paragraph.text.strip()

            if text:
                sections.append({
                    "location": f"Paragraph {paragraph_number}",
                    "paragraph": paragraph_number,
                    "text": text,
                })

        # Also extract text from tables, which often contain contract details.
        for table_number, table in enumerate(document.tables, start=1):
            for row_number, row in enumerate(table.rows, start=1):
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                row_text = " | ".join(
                    cell for cell in cells if cell
                )

                if row_text:
                    sections.append({
                        "location": (
                            f"Table {table_number}, row {row_number}"
                        ),
                        "table": table_number,
                        "row": row_number,
                        "text": row_text,
                    })

    if not sections:
        raise ValueError(
            "No readable text was found. "
            "The document may be empty or scanned."
        )

    return {
        "filename": path.name,
        "file_type": extension.lstrip("."),
        "section_count": len(sections),
        "sections": sections,
        "full_text": "\n\n".join(
            section["text"] for section in sections
        ),
    }
